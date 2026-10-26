"""Tests de sources/structures/recuperer_structure.py (sans réseau, adresses synthétiques)."""

import datetime as dt
import importlib.util
from pathlib import Path

import pytest

SPEC = importlib.util.spec_from_file_location(
    "recuperer_structure", Path(__file__).resolve().parent / "recuperer_structure.py")
rs = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(rs)

BASE = "http://exemple.invalide/intra"
URL = "http://exemple.invalide/intra/files/2026/10/Structure-09102026.xlsx"


def test_modele():
    assert rs.modele(URL) == ("http://exemple.invalide/intra", "Structure-", ".xlsx")
    with pytest.raises(ValueError):
        rs.modele("http://exemple.invalide/Structure.xlsx")


def test_adresses_mois_du_jour_et_suivant():
    a = rs.adresses("http://h", "Structure-", ".xlsx", dt.date(2026, 12, 30))
    assert a == ["http://h/files/2026/12/Structure-30122026.xlsx", "http://h/files/2027/01/Structure-30122026.xlsx"]


def test_derniere_version_trouvee():
    presentes = {"http://exemple.invalide/intra/files/2026/10/Structure-20102026.xlsx", URL}
    trouvee = rs.chercher_derniere(BASE, "Structure-", ".xlsx", dt.date(2026, 10, 26), 30, sonde=presentes.__contains__)
    assert trouvee.endswith("Structure-20102026.xlsx")


def test_aucune_version():
    assert rs.chercher_derniere(BASE, "Structure-", ".xlsx", dt.date(2026, 10, 26), 5, sonde=lambda u: False) is None


def test_fichier_deja_present(tmp_path):
    (tmp_path / "Structure-09102026.xlsx").write_bytes(b"x")
    assert rs.telecharger(URL, tmp_path).read_bytes() == b"x"


def test_copie_texte_sans_colonnes_de_personnes(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Structure UM"
    ws.cell(2, 2, "FICHIER STRUCTURE - titre")
    entete = ["Code Service", "Libellé Service", "Nom Chef Service\n", "Prénom Chef Service", "Code UF",
              "Nom Resp. UF", "Prénom Resp. UF", "Courriel"]
    for j, h in enumerate(entete, 3):
        ws.cell(4, j, h)
    for j, v in enumerate(["S1", "Médecine A | test", "Dupont", "Jean", "0001", "Martin", "Anne", "x@exemple.invalide"], 3):
        ws.cell(5, j, v)
    xlsx = tmp_path / "Structure-09102026.xlsx"
    wb.save(xlsx)
    fichiers, ecartees = rs.vers_texte(xlsx, tmp_path / "out")
    texte = fichiers[0].read_text(encoding="utf-8")
    assert fichiers[0].name == "structure-chu-2026-10-09__structure-um.txt"
    for interdit in ("Dupont", "Jean", "Martin", "Anne", "exemple.invalide", "titre"):
        assert interdit not in texte
    assert texte.splitlines()[0] == "| Code Service | Libellé Service | Code UF |"
    assert "Médecine A / test" in texte
    assert len(ecartees) == 5


def test_site_connu(monkeypatch):
    monkeypatch.delenv("STRUCTURE_SITE", raising=False)
    assert rs.site_connu(None, None) == rs.SITE
    assert rs.site_connu("http://h/s/", None) == "http://h/s"
    assert rs.site_connu(None, URL) == BASE
